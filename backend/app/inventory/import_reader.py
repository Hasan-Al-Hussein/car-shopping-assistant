"""Read-only workbook validation; no normalization, persistence or network work.

Only ``candidate_records`` from an accepted result may advance to BE-03. Original
cells remain internal provenance, separate from the safe serializable report.
"""

import hashlib
import math
import warnings
from dataclasses import asdict, dataclass
from dataclasses import field as dataclass_field
from io import BytesIO
from pathlib import Path
from typing import Any, Literal
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]
from pydantic import TypeAdapter, ValidationError

from app.api.schemas.common import InventoryRef
from app.api.schemas.inventory import ListingPhoto
from app.inventory.ooxml_structure import CellStorage, SheetStructure, inspect_structure

READER_VERSION = "workbook-reader-1"
CLEANED_HEADERS = (
    "Listing_ID",
    "year",
    "make",
    "model",
    "trim",
    "title",
    "description",
    "photo_url",
)
RAW_HEADERS = ("make", "model", "trim", "year", "title", "description", "photo_url")
MAX_FILE_BYTES = 16 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_ZIP_MEMBERS = 512
MAX_SCAN_ROWS = 20_000
MAX_DIAGNOSTICS = 2_000
# Frozen F-02 wire limits; originals exceeding these remain untouched provenance.
TEXT_LIMITS = {
    "make": 200,
    "model": 200,
    "trim": 200,
    "title": 1000,
    "description": 32000,
    "photo_url": 2048,
}
SOURCE_ID = TypeAdapter[str](InventoryRef.model_fields["source_id"].rebuild_annotation())
PARSER_ERRORS = (BadZipFile, ParseError, ValueError, KeyError, IndexError, TypeError, EOFError)


@dataclass(frozen=True)
class SourceSpec:
    """Trusted local-operator configuration, never a buyer request DTO."""

    namespace: str = "provided-cars-cleaned"
    selected_sheet: str = "cleaned dataset"
    audit_sheet: str = "raw dataset"
    source_version: str = "supplied-workbook-1"
    expected_sha256: str | None = None
    expected_selected_rows: int = 100
    expected_audit_rows: int = 100


@dataclass(frozen=True)
class SourceManifest:
    namespace: str
    workbook_sha256: str
    selected_sheet: str
    source_version: str
    reader_version: str = READER_VERSION


@dataclass(frozen=True)
class Diagnostic:
    """Codes and locators only: never interpolate source values or exceptions."""

    code: str
    severity: Literal["error", "warning"]
    sheet: str | None = None
    row: int | None = None
    cell: str | None = None
    field: str | None = None
    related_cells: tuple[str, ...] = ()


@dataclass(frozen=True)
class SourceCell:
    field: str
    coordinate: str
    value: object = dataclass_field(repr=False)
    data_type: str
    python_type: str
    storage: CellStorage | None = dataclass_field(default=None, repr=False)


@dataclass(frozen=True)
class SourceRecord:
    sheet: str
    row: int
    source_id: str | None = dataclass_field(repr=False)
    cells: tuple[SourceCell, ...] = dataclass_field(repr=False)
    structural_errors: tuple[str, ...] = ()

    def cell(self, name: str) -> SourceCell:
        return next(cell for cell in self.cells if cell.field == name)


@dataclass(frozen=True)
class SheetReport:
    sheet: str
    role: Literal["selected", "audit_only"]
    complete: bool
    meaningful_rows: int
    structurally_valid_rows: int
    structurally_invalid_rows: int
    structural_error_count: int
    diagnostics: tuple[Diagnostic, ...]
    omitted_diagnostics: int = 0


@dataclass(frozen=True)
class ImportReport:
    state: Literal["accepted", "rejected"]
    reader_version: str
    workbook_sha256_before: str | None
    workbook_sha256_after: str | None
    source_unchanged: bool
    accounting_complete: bool
    accepted_count: int
    quarantined_count: Literal[0]
    rejected_count: int
    sheets: tuple[SheetReport, ...]
    diagnostics: tuple[Diagnostic, ...]

    def to_dict(self) -> dict[str, Any]:
        """Safe diagnostics only; raw records must never be serialized as a report."""
        return asdict(self)


@dataclass(frozen=True)
class WorkbookReadResult:
    manifest: SourceManifest | None
    report: ImportReport
    selected_records: tuple[SourceRecord, ...] = dataclass_field(repr=False)
    audit_records: tuple[SourceRecord, ...] = dataclass_field(repr=False)

    @property
    def candidate_records(self) -> tuple[SourceRecord, ...]:
        return self.selected_records if self.report.state == "accepted" else ()


class _Diagnostics:
    def __init__(self) -> None:
        self.items: list[Diagnostic] = []
        self.total = 0
        self.errors = 0

    def add(self, diagnostic: Diagnostic) -> None:
        self.total += 1
        self.errors += diagnostic.severity == "error"
        if len(self.items) < MAX_DIAGNOSTICS:
            self.items.append(diagnostic)


def _populated(value: object) -> bool:
    # Whitespace, false, zero and formulas are populated source values.
    return value is not None and value != ""


def _source_id(cell: SourceCell) -> str | None:
    value = cell.value
    if cell.data_type in {"f", "e"} or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if abs(value) > 2**53 - 1 or not _integral(value):
            return None
        text = str(int(value))
    elif isinstance(value, str):
        text = value
    else:
        return None
    try:
        return SOURCE_ID.validate_python(text)
    except ValidationError:
        return None


def _integral(value: int | float) -> bool:
    return isinstance(value, int) or math.isfinite(value) and value.is_integer()


def _fact_issue(cell: SourceCell) -> str | None:
    value = cell.value
    if value is None or isinstance(value, str) and not value.strip():
        return "FACT_NOT_STATED"
    if cell.data_type in {"f", "e"} or isinstance(value, bool):
        return "FACT_UNSUPPORTED_TYPE"
    if isinstance(value, str) and len(value) > TEXT_LIMITS.get(cell.field, 32000):
        return "FACT_TOO_LONG"
    if cell.field == "year":
        if not isinstance(value, (int, float)):
            return "FACT_UNSUPPORTED_TYPE"
        if not _integral(value) or not 1000 <= value <= 9999:
            return "FACT_INVALID_YEAR"
    elif cell.field in {"model", "trim"} and isinstance(value, (int, float)):
        if not _integral(value):
            return "FACT_UNSUPPORTED_TYPE"
    elif not isinstance(value, str):
        return "FACT_UNSUPPORTED_TYPE"
    if cell.field == "photo_url" and isinstance(value, str):
        try:
            ListingPhoto(
                state="source_present", url=value, alt="Photo supplied with this car listing"
            )
        except ValidationError:
            return "PHOTO_POLICY_REJECTED"
    return None


def _inspect_sheet(
    workbook: Any,
    sheet_name: str,
    *,
    selected: bool,
    expected_count: int,
    structure: SheetStructure,
) -> tuple[SheetReport, tuple[SourceRecord, ...]]:
    diagnostics = _Diagnostics()
    records: list[SourceRecord] = []
    expected = CLEANED_HEADERS if selected else RAW_HEADERS
    headers: dict[str, int] = {}
    ids: dict[str, tuple[int, str]] = {}
    invalid_rows: set[int] = set()
    complete = structure.error is None

    def issue(
        code: str,
        *,
        severity: Literal["error", "warning"] = "error",
        row: int | None = None,
        column: int | None = None,
        name: str | None = None,
        related: tuple[str, ...] = (),
    ) -> None:
        coordinate = f"{get_column_letter(column)}{row}" if column and row else None
        diagnostics.add(Diagnostic(code, severity, sheet_name, row, coordinate, name, related))
        if severity == "error" and row is not None and row > 1:
            invalid_rows.add(row)

    if structure.error:
        issue(structure.error)
    if sheet_name not in workbook.sheetnames:
        issue("REQUIRED_SHEET_MISSING")
        complete = False
    elif sheet_name not in {sheet.title for sheet in workbook.worksheets}:
        issue("SHEET_NOT_WORKSHEET")
        complete = False
    else:
        sheet = workbook[sheet_name]
        # Do not trust XLSX dimension metadata: it may hide populated rows/cells.
        sheet.reset_dimensions()
        try:
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                if row_number > MAX_SCAN_ROWS:
                    issue("SHEET_SCAN_LIMIT")
                    complete = False
                    break
                if row_number == 1:
                    for column, cell in enumerate(row, start=1):
                        value = cell.value
                        if not _populated(value) or isinstance(value, str) and not value.strip():
                            continue
                        if cell.data_type in {"f", "e"} or not isinstance(value, str):
                            issue("HEADER_INVALID_TYPE", row=1, column=column)
                        elif value not in expected:
                            issue("HEADER_UNMAPPED", row=1, column=column)
                        elif value in headers:
                            issue(
                                "HEADER_DUPLICATE",
                                row=1,
                                column=column,
                                name=value,
                                related=(f"{get_column_letter(headers[value])}1",),
                            )
                        else:
                            headers[value] = column
                    for name in expected:
                        if name not in headers:
                            issue("HEADER_REQUIRED_MISSING", row=1, name=name)
                    continue
                if not any(_populated(cell.value) for cell in row):
                    continue
                row_codes: list[str] = []
                for column, cell in enumerate(row, start=1):
                    if _populated(cell.value) and column not in headers.values():
                        issue("POPULATED_COLUMN_UNMAPPED", row=row_number, column=column)
                        row_codes.append("POPULATED_COLUMN_UNMAPPED")
                cells = tuple(
                    SourceCell(
                        name,
                        f"{get_column_letter(column)}{row_number}",
                        row[column - 1].value if column <= len(row) else None,
                        row[column - 1].data_type if column <= len(row) else "n",
                        type(row[column - 1].value).__name__ if column <= len(row) else "NoneType",
                        structure.cells.get(f"{get_column_letter(column)}{row_number}"),
                    )
                    for name, column in headers.items()
                )
                identifier = None
                if selected:
                    id_cell = next((cell for cell in cells if cell.field == "Listing_ID"), None)
                    identifier = _source_id(id_cell) if id_cell else None
                    if identifier is None:
                        code = (
                            "SOURCE_ID_MISSING"
                            if id_cell is None or not _populated(id_cell.value)
                            else "SOURCE_ID_INVALID"
                        )
                        issue(
                            code,
                            row=row_number,
                            column=headers.get("Listing_ID"),
                            name="Listing_ID",
                        )
                        row_codes.append(code)
                    elif identifier in ids:
                        assert id_cell is not None
                        issue(
                            "SOURCE_ID_DUPLICATE",
                            row=row_number,
                            column=headers["Listing_ID"],
                            name="Listing_ID",
                            related=(ids[identifier][1],),
                        )
                        invalid_rows.add(ids[identifier][0])
                        row_codes.append("SOURCE_ID_DUPLICATE")
                    else:
                        assert id_cell is not None
                        ids[identifier] = (row_number, id_cell.coordinate)
                for cell in cells:
                    if cell.field == "Listing_ID":
                        continue
                    fact_issue = _fact_issue(cell)
                    if fact_issue:
                        issue(
                            fact_issue,
                            severity="warning",
                            row=row_number,
                            column=headers[cell.field],
                            name=cell.field,
                        )
                records.append(
                    SourceRecord(sheet_name, row_number, identifier, cells, tuple(row_codes))
                )
        except PARSER_ERRORS:
            issue("SHEET_UNREADABLE")
            complete = False
        if not headers:
            issue("HEADER_ROW_MISSING")
        if complete and len(records) != expected_count:
            issue("ROW_COUNT_MISMATCH")

    report = SheetReport(
        sheet=sheet_name,
        role="selected" if selected else "audit_only",
        complete=complete,
        meaningful_rows=len(records),
        structurally_valid_rows=len(records) - len(invalid_rows),
        structurally_invalid_rows=len(invalid_rows),
        structural_error_count=diagnostics.errors,
        diagnostics=tuple(diagnostics.items),
        omitted_diagnostics=diagnostics.total - len(diagnostics.items),
    )
    return report, tuple(records)


def _read_bytes(path: Path) -> bytes:
    with path.open("rb") as stream:
        data = stream.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("WORKBOOK_SIZE_LIMIT")
    return data


def read_workbook(path: Path, spec: SourceSpec | None = None) -> WorkbookReadResult:
    """Inspect a configured local source; raw-only defects do not merge into scope.

    Global decode/integrity failures or selected-sheet structural failures reject
    the entire candidate. Independent raw-sheet problems remain audit diagnostics.
    File/ZIP/row limits fail closed and explicitly mark incomplete accounting.
    """
    spec = spec if spec is not None else SourceSpec()
    global_issues: list[Diagnostic] = []
    sheets: list[SheetReport] = []
    selected_records: tuple[SourceRecord, ...] = ()
    audit_records: tuple[SourceRecord, ...] = ()
    before = after = None
    manifest = None
    try:
        data = _read_bytes(path)
        before = hashlib.sha256(data).hexdigest()
        manifest = SourceManifest(spec.namespace, before, spec.selected_sheet, spec.source_version)
        if spec.expected_sha256 is not None and before != spec.expected_sha256:
            global_issues.append(Diagnostic("SOURCE_HASH_MISMATCH", "error"))
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if (
                len(members) > MAX_ZIP_MEMBERS
                or sum(m.file_size for m in members) > MAX_EXPANDED_BYTES
            ):
                raise ValueError("WORKBOOK_ARCHIVE_LIMIT")
            if len({member.filename for member in members}) != len(members):
                raise ValueError("WORKBOOK_DUPLICATE_PART")
            structure = inspect_structure(
                archive, (spec.selected_sheet, spec.audit_sheet), max_rows=MAX_SCAN_ROWS
            )
        with warnings.catch_warnings(record=True) as captured_warnings:
            warnings.simplefilter("always")
            workbook = load_workbook(
                BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            try:
                selected, selected_records = _inspect_sheet(
                    workbook,
                    spec.selected_sheet,
                    selected=True,
                    expected_count=spec.expected_selected_rows,
                    structure=structure[spec.selected_sheet],
                )
                sheets.append(selected)
                audit, audit_records = _inspect_sheet(
                    workbook,
                    spec.audit_sheet,
                    selected=False,
                    expected_count=spec.expected_audit_rows,
                    structure=structure[spec.audit_sheet],
                )
                sheets.append(audit)
            finally:
                workbook.close()
            if captured_warnings:
                global_issues.append(Diagnostic("WORKBOOK_PARSER_WARNING", "warning"))
    except (OSError, InvalidFileException, *PARSER_ERRORS):
        global_issues.append(Diagnostic("WORKBOOK_UNREADABLE_OR_LIMIT", "error"))
    if before is not None:
        try:
            after = hashlib.sha256(_read_bytes(path)).hexdigest()
        except (OSError, ValueError):
            global_issues.append(Diagnostic("SOURCE_RECHECK_FAILED", "error"))
        if after != before:
            global_issues.append(Diagnostic("SOURCE_CHANGED_DURING_READ", "error"))
    selected_report = next((sheet for sheet in sheets if sheet.role == "selected"), None)
    accepted = (
        not any(issue.severity == "error" for issue in global_issues)
        and selected_report is not None
        and selected_report.complete
        and selected_report.structural_error_count == 0
    )
    selected_count = selected_report.meaningful_rows if selected_report else 0
    report = ImportReport(
        state="accepted" if accepted else "rejected",
        reader_version=READER_VERSION,
        workbook_sha256_before=before,
        workbook_sha256_after=after,
        source_unchanged=before is not None and before == after,
        accounting_complete=len(sheets) == 2 and all(sheet.complete for sheet in sheets),
        accepted_count=selected_count if accepted else 0,
        quarantined_count=0,
        rejected_count=0 if accepted else selected_count,
        sheets=tuple(sheets),
        diagnostics=tuple(global_issues),
    )
    return WorkbookReadResult(manifest, report, selected_records, audit_records)
