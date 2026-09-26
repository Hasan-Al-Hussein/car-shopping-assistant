"""Real-source and malformed-copy acceptance cases for the read-only reader."""

import hashlib
import json
import socket
import warnings
from collections.abc import Callable
from dataclasses import replace
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import Element, SubElement, fromstring, tostring
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.chart import BarChart  # type: ignore[import-untyped]
from openpyxl.workbook.workbook import Workbook  # type: ignore[import-untyped]

from app.inventory import import_reader
from app.inventory.import_reader import SourceSpec, WorkbookReadResult, read_workbook
from app.inventory.ooxml_structure import SHEET_NS
from tests.support.harness import PROJECT


@pytest.fixture(scope="module")
def source_facts() -> dict[str, Any]:
    parsed = json.loads((PROJECT / "fixtures/shared/source-facts.json").read_text(encoding="utf-8"))
    assert isinstance(parsed, dict)
    return parsed


@pytest.fixture(scope="module")
def source_path(source_facts: dict[str, Any]) -> Path:
    relative_path = source_facts["source_relative_path"]
    assert isinstance(relative_path, str)
    return PROJECT / relative_path


@pytest.fixture(scope="module")
def source_bytes(source_path: Path) -> bytes:
    return source_path.read_bytes()


@pytest.fixture
def make_copy(tmp_path: Path, source_bytes: bytes) -> Callable[[Callable[[Any], None]], Path]:
    def copied(change: Callable[[Any], None]) -> Path:
        workbook = load_workbook(BytesIO(source_bytes), data_only=False, keep_links=False)
        try:
            change(workbook)
            destination = tmp_path / "malformed-source-copy.xlsx"
            workbook.save(destination)
            return destination
        finally:
            workbook.close()

    return copied


def codes(result: WorkbookReadResult, *, selected: bool = True) -> set[str]:
    role = "selected" if selected else "audit_only"
    return {
        diagnostic.code
        for sheet in result.report.sheets
        if sheet.role == role
        for diagnostic in sheet.diagnostics
    }


def rewrite_xml(
    source: bytes, destination: Path, part: str, mutate: Callable[[Element], None]
) -> Path:
    with ZipFile(BytesIO(source)) as original, ZipFile(destination, "w") as copied:
        for info in original.infolist():
            payload = original.read(info.filename)
            if info.filename == part:
                root = fromstring(payload)
                mutate(root)
                payload = tostring(root, encoding="utf-8", xml_declaration=True)
            copied.writestr(info, payload)
    return destination


def test_real_source_is_unchanged_and_raw_is_separate(
    source_path: Path,
    source_bytes: bytes,
    source_facts: dict[str, Any],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def forbidden(*args: object, **kwargs: object) -> None:
        raise AssertionError("READER_MUST_NOT_SAVE_OR_FETCH")

    monkeypatch.setattr(Workbook, "save", forbidden)
    monkeypatch.setattr(socket, "create_connection", forbidden)
    result = read_workbook(source_path, SourceSpec(expected_sha256=source_facts["workbook_sha256"]))
    report = result.report
    assert report.state == "accepted"
    assert (report.accepted_count, report.rejected_count, report.quarantined_count) == (100, 0, 0)
    assert report.accounting_complete and report.source_unchanged
    assert (
        report.workbook_sha256_before
        == report.workbook_sha256_after
        == source_facts["workbook_sha256"]
    )
    assert source_path.read_bytes() == source_bytes
    assert len(result.candidate_records) == len(result.audit_records) == 100
    assert [sheet.meaningful_rows for sheet in report.sheets] == [100, 100]
    assert {record.source_id for record in result.candidate_records} == {
        str(i) for i in range(1, 101)
    }
    assert all(record.source_id is None for record in result.audit_records)
    assert not report.diagnostics
    assert all(sheet.structural_error_count == 0 for sheet in report.sheets)


def test_exact_shared_source_fixture_provenance(
    source_path: Path, source_facts: dict[str, Any]
) -> None:
    result = read_workbook(source_path)
    records = {record.source_id: record for record in result.candidate_records}
    for case in source_facts["cases"]:
        record = records[case["source_id"]]
        cell = next(cell for cell in record.cells if cell.coordinate == case["cell"])
        assert record.sheet == case["sheet"]
        assert cell.value == case["raw_value"]
        assert cell.python_type == case["raw_type"]
    # Raw G62 shares a URL with cleaned H60, but has no adopted listing ID.
    assert records["59"].cell("photo_url").coordinate == "H60"
    raw = next(record for record in result.audit_records if record.row == 62)
    assert raw.cell("photo_url").coordinate == "G62"
    assert raw.cell("photo_url").value == records["59"].cell("photo_url").value
    assert raw.source_id is None


@pytest.mark.parametrize(
    "value,expected",
    [
        (None, "HEADER_REQUIRED_MISSING"),
        ("unexpected", "HEADER_UNMAPPED"),
        (123, "HEADER_INVALID_TYPE"),
        ("=1+1", "HEADER_INVALID_TYPE"),
        ("make", "HEADER_DUPLICATE"),
    ],
)
def test_bad_headers_reject_whole_candidate(
    make_copy: Callable[[Callable[[Any], None]], Path],
    value: object,
    expected: str,
) -> None:
    path = make_copy(lambda book: setattr(book["cleaned dataset"]["D1"], "value", value))
    result = read_workbook(path)
    assert result.report.state == "rejected"
    assert result.report.rejected_count == 100
    assert expected in codes(result)
    assert result.candidate_records == ()


def test_reordered_headers_keep_actual_cell_mapping(
    make_copy: Callable[[Callable[[Any], None]], Path],
) -> None:
    def reorder(book: Any) -> None:
        sheet = book["cleaned dataset"]
        for row in range(1, 102):
            sheet.cell(row, 1).value, sheet.cell(row, 8).value = (
                sheet.cell(row, 8).value,
                sheet.cell(row, 1).value,
            )

    result = read_workbook(make_copy(reorder))
    assert result.report.state == "accepted"
    record = next(item for item in result.candidate_records if item.source_id == "59")
    assert record.cell("Listing_ID").coordinate == "H60"
    assert record.cell("photo_url").coordinate == "A60"


@pytest.mark.parametrize("header", [None, "unreviewed column"])
def test_unmapped_populated_column_rejects_even_without_header(
    make_copy: Callable[[Callable[[Any], None]], Path],
    header: str | None,
) -> None:
    def change(book: Any) -> None:
        book["cleaned dataset"]["K1"] = header
        book["cleaned dataset"]["K2"] = "synthetic unmapped value"

    result = read_workbook(make_copy(change))
    assert "POPULATED_COLUMN_UNMAPPED" in codes(result)
    assert result.report.state == "rejected"
    assert result.report.accepted_count == 0
    assert result.report.rejected_count == 100


@pytest.mark.parametrize("value", [None, "", " ", True, 1.5, "=1+1", datetime(2020, 1, 1)])
def test_bad_ids_never_generate_a_candidate(
    make_copy: Callable[[Callable[[Any], None]], Path],
    value: object,
) -> None:
    result = read_workbook(
        make_copy(lambda book: setattr(book["cleaned dataset"]["A2"], "value", value))
    )
    assert result.report.state == "rejected"
    assert codes(result) & {"SOURCE_ID_MISSING", "SOURCE_ID_INVALID"}
    assert result.report.sheets[0].meaningful_rows == 100
    assert result.candidate_records == ()


def test_duplicate_id_numeric_string_collision_counts_both_rows(
    make_copy: Callable[[Callable[[Any], None]], Path],
) -> None:
    result = read_workbook(
        make_copy(lambda book: setattr(book["cleaned dataset"]["A3"], "value", "1"))
    )
    selected = result.report.sheets[0]
    assert result.report.state == "rejected"
    assert (selected.structurally_valid_rows, selected.structurally_invalid_rows) == (98, 2)
    duplicate = next(issue for issue in selected.diagnostics if issue.code == "SOURCE_ID_DUPLICATE")
    assert duplicate.cell == "A3" and duplicate.related_cells == ("A2",)
    assert (result.report.accepted_count, result.report.rejected_count) == (0, 100)


def test_textual_id_representation_is_preserved(
    make_copy: Callable[[Callable[[Any], None]], Path],
) -> None:
    result = read_workbook(
        make_copy(lambda book: setattr(book["cleaned dataset"]["A3"], "value", "001"))
    )
    assert result.report.state == "accepted"
    assert result.candidate_records[1].source_id == "001"
    assert result.candidate_records[1].cell("Listing_ID").value == "001"


@pytest.mark.parametrize(
    "cell,value",
    [
        ("B2", True),
        ("B2", 2020.5),
        ("B2", "unknown"),
        ("G2", "=1+1"),
        ("G2", "#VALUE!"),
        ("G2", datetime(2020, 1, 1)),
        ("F2", None),
        ("H2", "https://example.invalid/private?token=secret"),
        ("H2", None),
    ],
)
def test_invalid_facts_and_photos_retain_the_listing(
    make_copy: Callable[[Callable[[Any], None]], Path],
    cell: str,
    value: object,
) -> None:
    result = read_workbook(
        make_copy(lambda book: setattr(book["cleaned dataset"][cell], "value", value))
    )
    assert result.report.state == "accepted"
    assert result.report.accepted_count == 100
    assert any(
        issue.cell == cell and issue.severity == "warning"
        for issue in result.report.sheets[0].diagnostics
    )
    source_cell = next(
        item for item in result.candidate_records[0].cells if item.coordinate == cell
    )
    assert source_cell.value == value


def test_formatting_tails_and_blank_gaps_do_not_hide_nonempty_rows(
    make_copy: Callable[[Callable[[Any], None]], Path],
) -> None:
    def change(book: Any) -> None:
        sheet = book["cleaned dataset"]
        sheet["Z1000"].number_format = "@"
        sheet["A105"] = "extra"

    result = read_workbook(make_copy(change), SourceSpec(expected_selected_rows=101))
    assert result.report.state == "accepted"
    assert result.report.accepted_count == 101
    assert result.candidate_records[-1].row == 105
    assert result.candidate_records[-1].source_id == "extra"


def test_unmapped_only_row_is_counted_and_rejected(
    make_copy: Callable[[Callable[[Any], None]], Path],
) -> None:
    result = read_workbook(
        make_copy(lambda book: setattr(book["cleaned dataset"]["K105"], "value", "only here"))
    )
    assert result.report.state == "rejected"
    assert result.report.rejected_count == 101
    assert {"POPULATED_COLUMN_UNMAPPED", "SOURCE_ID_MISSING", "ROW_COUNT_MISMATCH"} <= codes(result)


@pytest.mark.parametrize("dimension", ["A1:A1", "A1:XFD1048576"])
def test_untrusted_dimension_metadata_cannot_hide_or_invent_rows(
    source_bytes: bytes,
    tmp_path: Path,
    dimension: str,
) -> None:
    def change(root: Element) -> None:
        elements = root.findall(f"{{{SHEET_NS}}}dimension")
        assert len(elements) <= 1
        if elements:
            element = elements[0]
        else:
            element = Element(f"{{{SHEET_NS}}}dimension")
            root.insert(0, element)
        element.set("ref", dimension)

    destination = rewrite_xml(
        source_bytes, tmp_path / "dimension-copy.xlsx", "xl/worksheets/sheet2.xml", change
    )
    result = read_workbook(destination)
    assert result.report.state == "accepted"
    assert result.report.accepted_count == 100


def test_raw_only_error_is_independent_of_cleaned_scope(
    make_copy: Callable[[Callable[[Any], None]], Path],
) -> None:
    result = read_workbook(
        make_copy(lambda book: setattr(book["raw dataset"]["J2"], "value", "audit issue"))
    )
    assert result.report.state == "accepted"
    assert "POPULATED_COLUMN_UNMAPPED" in codes(result, selected=False)
    assert result.report.sheets[1].structural_error_count > 0
    assert len(result.candidate_records) == len(result.audit_records) == 100


def test_safe_report_excludes_injected_source_text(
    make_copy: Callable[[Callable[[Any], None]], Path],
) -> None:
    secret = "synthetic-contact-555-token-DO-NOT-LOG"

    def change(book: Any) -> None:
        sheet = book["cleaned dataset"]
        sheet["A2"] = secret + "!"
        sheet["J1"] = secret
        sheet["G3"] = "=" + secret
        sheet["H4"] = "https://example.invalid/" + secret

    result = read_workbook(make_copy(change))
    assert result.report.state == "rejected"
    assert secret not in json.dumps(result.report.to_dict())
    assert secret not in repr(result)
    assert secret not in repr(result.selected_records)


def test_missing_sheet_and_corrupt_workbook_fail_closed(
    make_copy: Callable[[Callable[[Any], None]], Path],
    tmp_path: Path,
) -> None:
    def remove_selected(book: Any) -> None:
        book.remove(book["cleaned dataset"])
        book["raw dataset"].sheet_state = "visible"

    missing = read_workbook(make_copy(remove_selected))
    assert missing.report.state == "rejected" and not missing.report.accounting_complete
    assert "REQUIRED_SHEET_MISSING" in codes(missing)
    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_bytes(b"not a workbook - synthetic private text")
    broken = read_workbook(corrupt)
    assert broken.report.state == "rejected" and not broken.report.accounting_complete
    assert broken.candidate_records == ()
    assert "synthetic private text" not in json.dumps(broken.report.to_dict())


def test_source_hash_mismatch_and_changed_source_reject_all_rows(
    source_path: Path,
    source_bytes: bytes,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    wrong = read_workbook(source_path, replace(SourceSpec(), expected_sha256="0" * 64))
    assert wrong.report.state == "rejected" and wrong.report.rejected_count == 100
    assert wrong.report.source_unchanged
    assert wrong.report.diagnostics[0].code == "SOURCE_HASH_MISMATCH"
    reads = iter((source_bytes, b"replacement source"))
    monkeypatch.setattr(import_reader, "_read_bytes", lambda path: next(reads))
    changed = read_workbook(source_path)
    assert changed.report.state == "rejected" and not changed.report.source_unchanged
    assert changed.candidate_records == ()
    assert changed.report.workbook_sha256_before == hashlib.sha256(source_bytes).hexdigest()


def test_scan_limit_reports_incomplete_accounting(
    source_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(import_reader, "MAX_SCAN_ROWS", 3)
    result = read_workbook(source_path)
    assert result.report.state == "rejected"
    assert not result.report.accounting_complete
    assert "SHEET_SCAN_LIMIT" in codes(result)
    assert result.candidate_records == ()


@pytest.mark.parametrize(
    "case",
    [
        "out_of_order_cell",
        "duplicate_cell",
        "duplicate_row",
        "wrong_parent_row",
        "nested_cell",
        "mismatched_storage",
    ],
)
def test_malformed_xml_coordinates_cannot_be_silently_discarded(
    source_bytes: bytes,
    tmp_path: Path,
    case: str,
) -> None:
    def mutate(root: Element) -> None:
        data = root.find(f"{{{SHEET_NS}}}sheetData")
        assert data is not None
        row = next(item for item in data if item.get("r") == "2")
        if case == "duplicate_row":
            row = SubElement(data, f"{{{SHEET_NS}}}row", {"r": "2"})
            coordinate = "A2"
        elif case == "duplicate_cell":
            coordinate = "H2"
        elif case == "wrong_parent_row":
            coordinate = "K3"
        elif case == "nested_cell":
            row = next(item for item in row if item.get("r") == "H2")
            coordinate = "I2"
        else:
            coordinate = "K2"
        cell = Element(f"{{{SHEET_NS}}}c", {"r": coordinate, "t": "inlineStr"})
        if case == "mismatched_storage":
            SubElement(cell, f"{{{SHEET_NS}}}v").text = "synthetic hidden cell"
        else:
            SubElement(
                SubElement(cell, f"{{{SHEET_NS}}}is"), f"{{{SHEET_NS}}}t"
            ).text = "synthetic hidden cell"
        if case == "out_of_order_cell":
            row.insert(len(row) - 1, cell)
        else:
            row.append(cell)

    path = rewrite_xml(
        source_bytes, tmp_path / "bad-coordinates.xlsx", "xl/worksheets/sheet2.xml", mutate
    )
    result = read_workbook(path)
    assert result.report.state == "rejected"
    assert not result.report.accounting_complete
    assert result.candidate_records == ()
    assert codes(result) & {
        "CELL_COORDINATE_INVALID",
        "CELL_COORDINATE_DUPLICATE_OR_UNORDERED",
        "ROW_COORDINATE_DUPLICATE_OR_UNORDERED",
        "WORKSHEET_HIERARCHY_INVALID",
        "CELL_STORAGE_INVALID",
    }


def test_library_warning_never_exposes_defined_name_text(
    source_bytes: bytes,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    marker = "synthetic_contact_token_DO_NOT_LOG"

    def mutate(root: Element) -> None:
        names = root.find(f"{{{SHEET_NS}}}definedNames")
        if names is None:
            names = SubElement(root, f"{{{SHEET_NS}}}definedNames")
        SubElement(
            names, f"{{{SHEET_NS}}}definedName", {"name": "_xlnm.Print_Area", "localSheetId": "1"}
        ).text = marker

    path = rewrite_xml(source_bytes, tmp_path / "warning-copy.xlsx", "xl/workbook.xml", mutate)
    with warnings.catch_warnings(record=True) as leaked:
        warnings.simplefilter("always")
        result = read_workbook(path)
    assert result.report.state == "accepted"
    assert "WORKBOOK_PARSER_WARNING" in {issue.code for issue in result.report.diagnostics}
    captured = capsys.readouterr()
    assert marker not in captured.out + captured.err + caplog.text + json.dumps(
        result.report.to_dict()
    )
    assert not leaked


def test_invalid_date_serial_retains_storage_token_without_warning_leak(
    make_copy: Callable[[Callable[[Any], None]], Path],
    capsys: pytest.CaptureFixture[str],
) -> None:
    def change(book: Any) -> None:
        book["cleaned dataset"]["B2"] = 15551234567
        book["cleaned dataset"]["B2"].number_format = "mm/dd/yyyy"

    path = make_copy(change)
    with warnings.catch_warnings(record=True) as leaked:
        warnings.simplefilter("always")
        result = read_workbook(path)
    assert result.report.state == "accepted"
    cell = result.candidate_records[0].cell("year")
    assert cell.storage is not None and cell.storage.value == "15551234567"
    assert "FACT_UNSUPPORTED_TYPE" in codes(result)
    captured = capsys.readouterr()
    assert "15551234567" not in captured.out + captured.err
    assert not leaked


def test_chartsheet_named_as_selected_source_is_rejected(
    make_copy: Callable[[Callable[[Any], None]], Path],
) -> None:
    def change(book: Any) -> None:
        book.remove(book["cleaned dataset"])
        chart_sheet = book.create_chartsheet("cleaned dataset")
        chart_sheet.add_chart(BarChart())

    result = read_workbook(make_copy(change))
    assert result.report.state == "rejected"
    assert "SHEET_NOT_WORKSHEET" in codes(result)
    assert result.candidate_records == ()


@pytest.mark.parametrize(
    "cell,limit", [("C2", 200), ("D2", 200), ("E2", 200), ("F2", 1000), ("G2", 32000)]
)
def test_field_length_boundaries_keep_original_values(
    make_copy: Callable[[Callable[[Any], None]], Path],
    cell: str,
    limit: int,
) -> None:
    at_limit = read_workbook(
        make_copy(lambda book: setattr(book["cleaned dataset"][cell], "value", "x" * limit))
    )
    assert at_limit.report.state == "accepted"
    assert not any(
        issue.code == "FACT_TOO_LONG" and issue.cell == cell
        for issue in at_limit.report.sheets[0].diagnostics
    )
    beyond = read_workbook(
        make_copy(lambda book: setattr(book["cleaned dataset"][cell], "value", "x" * (limit + 1)))
    )
    assert beyond.report.state == "accepted"
    assert any(
        issue.code == "FACT_TOO_LONG" and issue.cell == cell
        for issue in beyond.report.sheets[0].diagnostics
    )
    original = next(item for item in beyond.candidate_records[0].cells if item.coordinate == cell)
    assert original.value == "x" * (limit + 1)


def test_file_size_limit_is_fail_closed(source_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(import_reader, "MAX_FILE_BYTES", 10)
    result = read_workbook(source_path)
    assert result.report.state == "rejected" and not result.report.accounting_complete
    assert result.candidate_records == ()
